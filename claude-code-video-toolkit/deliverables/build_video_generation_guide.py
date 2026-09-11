from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables' / '财经双人对话视频生成指令手册.docx'
ASSETS = ROOT / 'deliverables' / 'video-generation-guide-assets'

NAVY = '0A1724'
INK = '142233'
MUTED = '64748B'
RED = 'E94B51'
CYAN = '22B8CF'
GOLD = 'EAB84D'
PALE_BLUE = 'E8EEF5'
PALE_RED = 'FCECEF'
PALE_CYAN = 'EAF8FB'
PALE_GOLD = 'FFF6DD'
LIGHT = 'F6F8FB'
WHITE = 'FFFFFF'


CHINESE_FONT = 'Arial Unicode MS'


def set_run_font(run, *, size=None, color=None, bold=None, italic=None, name=CHINESE_FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run._element.rPr.rFonts.set(qn('w:hint'), 'eastAsia')
    lang = run._element.rPr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        run._element.rPr.append(lang)
    lang.set(qn('w:val'), 'zh-CN')
    lang.set(qn('w:eastAsia'), 'zh-CN')
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(element, fill):
    p_pr = element.get_or_add_pPr() if hasattr(element, 'get_or_add_pPr') else element
    shd = p_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        p_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def cell_shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    tc_w.set(qn('w:w'), str(dxa))
    tc_w.set(qn('w:type'), 'dxa')
    cell.width = Inches(dxa / 1440)


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.first_child_found_in('w:tblLayout')
    if tbl_layout is None:
        tbl_layout = OxmlElement('w:tblLayout')
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn('w:type'), 'fixed')
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(total))
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_ind = tbl_pr.first_child_found_in('w:tblInd')
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), '120')
    tbl_ind.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for idx, width in enumerate(widths):
        if idx < len(grid.gridCol_lst):
            grid.gridCol_lst[idx].set(qn('w:w'), str(width))
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_border(cell, *, color='D9E1EA', size='6'):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in('w:tcBorders')
    if tc_borders is None:
        tc_borders = OxmlElement('w:tcBorders')
        tc_pr.append(tc_borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = qn(f'w:{edge}')
        edge_el = tc_borders.find(tag)
        if edge_el is None:
            edge_el = OxmlElement(f'w:{edge}')
            tc_borders.append(edge_el)
        edge_el.set(qn('w:val'), 'single')
        edge_el.set(qn('w:sz'), size)
        edge_el.set(qn('w:space'), '0')
        edge_el.set(qn('w:color'), color)


def set_paragraph_border(paragraph, color=RED, size='14', space='6'):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn('w:pBdr'))
    if borders is None:
        borders = OxmlElement('w:pBdr')
        p_pr.append(borders)
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), size)
    left.set(qn('w:space'), space)
    left.set(qn('w:color'), color)
    borders.append(left)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    node = OxmlElement('w:keepNext')
    p_pr.append(node)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, label, url):
    part = paragraph.part
    relation_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), relation_id)
    new_run = OxmlElement('w:r')
    r_pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), RED)
    r_pr.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    r_pr.append(underline)
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), CHINESE_FONT)
    fonts.set(qn('w:hAnsi'), CHINESE_FONT)
    fonts.set(qn('w:eastAsia'), CHINESE_FONT)
    fonts.set(qn('w:hint'), 'eastAsia')
    r_pr.append(fonts)
    new_run.append(r_pr)
    text = OxmlElement('w:t')
    text.text = label
    new_run.append(text)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def set_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = 'PAGE'
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_run_font(run, **kwargs)
    return run


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    add_text(p, text, size={1: 16, 2: 13, 3: 12}[level], color={1: '2E74B5', 2: '2E74B5', 3: '1F4D78'}[level], bold=True)
    return p


def add_body(doc, text='', *, bold_prefix=None, after=6, color=INK, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        add_text(p, bold_prefix, size=size, color=color, bold=True)
        add_text(p, text[len(bold_prefix):], size=size, color=color)
    else:
        add_text(p, text, size=size, color=color)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.left_indent = Inches(.375)
    p.paragraph_format.first_line_indent = Inches(-.188)
    add_text(p, text, size=11, color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.left_indent = Inches(.375)
    p.paragraph_format.first_line_indent = Inches(-.188)
    add_text(p, text, size=11, color=INK)
    return p


def add_callout(doc, label, text, fill=PALE_BLUE, accent=CYAN):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    cell_shade(cell, fill)
    set_cell_border(cell, color=accent, size='12')
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    add_text(p, label + '  ', size=10.5, color=accent, bold=True)
    add_text(p, text, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(doc, text):
    for index, line in enumerate(text.splitlines()):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0 if index else 2)
        p.paragraph_format.space_after = Pt(0 if index < len(text.splitlines()) - 1 else 6)
        p.paragraph_format.line_spacing = 1.12
        shade(p._p, 'F2F5F8')
        set_paragraph_border(p, color='C9D5E3', size='6', space='5')
        add_text(p, line or ' ', size=9.2, color='233142')


def add_image(doc, filename, caption, width=Inches(6.1)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(ASSETS / filename), width=width)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    add_text(cap, caption, size=9, color=MUTED, italic=True)


def add_asset_link(doc, title, path, use, duration):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    add_text(p, title + '｜', size=11, color=INK, bold=True)
    add_text(p, use + '（' + duration + '）', size=11, color=INK)
    link = doc.add_paragraph()
    link.paragraph_format.space_after = Pt(7)
    add_hyperlink(link, '点击打开视频文件', 'file://' + path)
    add_text(link, '\n' + path, size=8.2, color=MUTED, name='Menlo')


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        cell_shade(cell, PALE_BLUE)
        set_cell_border(cell, color='B7C7D9')
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, text, size=10, color=INK, bold=True)
    for row_data in rows:
        row = table.add_row()
        for i, text in enumerate(row_data):
            cell = row.cells[i]
            set_cell_border(cell, color='D7E0E9')
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            add_text(p, text, size=9.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(.492)
    section.footer_distance = Inches(.492)

    normal = doc.styles['Normal']
    normal.font.name = CHINESE_FONT
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), CHINESE_FONT)
    normal._element.rPr.rFonts.set(qn('w:hint'), 'eastAsia')
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, color, before, after in [(1, 16, '2E74B5', 18, 10), (2, 13, '2E74B5', 14, 7), (3, 12, '1F4D78', 10, 5)]:
        style = doc.styles[f'Heading {level}']
        style.font.name = CHINESE_FONT
        style._element.rPr.rFonts.set(qn('w:eastAsia'), CHINESE_FONT)
        style._element.rPr.rFonts.set(qn('w:hint'), 'eastAsia')
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    add_text(p, '《财报揭秘室》｜财经双人对话视频生成工作手册', size=8.5, color=MUTED)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(p, '内部制作参考  ·  ', size=8.5, color=MUTED)
    set_page_number(p)


def build():
    doc = Document()
    configure_document(doc)

    # Cover — editorial cover override within compact reference guide.
    for _ in range(5):
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    add_text(kicker, 'FINANCIAL REPORT / REVEALED', size=11, color=RED, bold=True)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    add_text(title, '财经双人对话视频\n可复用生成指令手册', size=29, color=NAVY, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    add_text(subtitle, '适用于《财报揭秘室》 · Remotion 动效 · 豆包 TTS', size=13, color=MUTED)
    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rule.paragraph_format.space_after = Pt(24)
    add_text(rule, '▁▁▁▁▁▁▁▁▁▁', size=16, color=RED)
    add_image(doc, '03_active-speaker.png', '成片核心观感：发言者高亮、关键句短卡、数据画面为主。', width=Inches(5.85))
    cover_note = doc.add_paragraph()
    cover_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_note.paragraph_format.space_before = Pt(9)
    add_text(cover_note, '版本：已验证制作流程  ·  输出：16:9 横版财经解释视频', size=9.5, color=MUTED)
    doc.add_page_break()

    add_heading(doc, '先读这一页：本片的“正确公式”', 1)
    add_callout(doc, '目标观感', '不是“念全文 + 静态 PPT”。而是：声音先行，数据/关键词随后出现；每次只突出一位主持人；章节与片尾都服务于节奏，不额外拖长。', PALE_RED, RED)
    add_table(doc, ['模块', '本次已验证的做法', '不要这样做'], [
        ('开头', '先用 1 个问题/反差数据做钩子，再接 2.23 秒栏目片头。', '开场先放品牌片头，或一开始把所有数据抛完。'),
        ('人物', '当前发言者 130px 高亮头像 + 一句关键结论；另一人 64px、低透明度。', '整段旁白塞进对话框，双方头像同等抢画面。'),
        ('章节', '上一段结束后的 1.4–1.6 秒，先展示章节标题；标题退出后再进下一位人物。', '小标题与人物对话同时出现，或空白停 1 秒。'),
        ('数据', '按照旁白中的数字、因果和验证点渐进出现。', '画面一进入就把全部卡片、折线和结论放完。'),
        ('收尾', '最后一句“每期三分钟…”起，全屏主标动效；女声结束即终止视频。', '男声仍在分析时提前上片尾，或片尾后再加尾音/空白。'),
    ], [1550, 5050, 2760])

    add_heading(doc, '一、下次先这样给 Codex 说（两步式指令）', 1)
    add_heading(doc, '步骤 A：先让它规划，不要直接渲染', 2)
    add_code_block(doc, '''我要做一支《财报揭秘室》风格的横版财经双人对话视频。请先不要生成视频，先阅读我提供的解读稿、人物头像、数据素材、音效和参考视频。

目标：约 [X] 分钟、1920×1080 / 30fps、16:9；不需要口型，只要头像与数据动效。
请输出：
1. 压缩后的完整男女对话文案（标明男女、每句的核心结论）；
2. 逐镜头分镜表：时间段、发言者、画面数据、关键句、音效、转场；
3. 每章节的标题卡位置与时长；
4. 片头/片尾如何接入现有动效；
5. 需要我确认的数字、音色和素材缺口。

节奏要求：钩子 → 片头 → 章节分析 → 全屏片尾；数据必须随旁白渐进出现；章节标题必须先出现并消失，再进入人物解说。''')
    add_heading(doc, '步骤 B：确认方案后，再让它执行', 2)
    add_code_block(doc, '''按已确认的分镜执行并渲染预览。使用豆包 TTS：男声“刘飞 2.0”、女声“流畅女声 2.0”。

人物规则：正在说话的人高亮放大；另一人缩小、弱化；对话卡只显示一句关键结论，不能直接放完整旁白。
音画规则：每条音频以真实时长写入 timeline；数据/卡片按句中信息分批出现；章节标题只占上一句结束后的转场空档；禁止标题和人物卡同屏。
片尾规则：最后女声“每期三分钟，拆解财报里的关键信号”开始时切入全屏主标动效，其他 UI 全部关闭；女声结束的同一帧结束视频，不加尾音、不留空白。

先输出 3 个关键时间点的静帧供我确认：一个人物切换、一个章节转场、一个片尾。确认后再导出完整 MP4。''')

    add_heading(doc, '二、必须给到的素材与信息', 1)
    add_table(doc, ['输入', '最小要求', '交付建议'], [
        ('解读稿', '带章节小标题，数字口径明确。', 'DOCX 或 Markdown；把“必须出现”的数字加粗或单列。'),
        ('头像', '男女主持人各 1 张，正面、清晰、背景干净。', 'PNG；人物脸部占画面约 50% 以上。'),
        ('品牌资产', '透明 Logo、主视觉色、免责声明。', 'PNG/SVG；注明是否允许在右上角常驻。'),
        ('声音', '男/女声选择、TTS 凭证、是否允许语速调整。', '先生成 2–4 个试听片段，确认后再批量生成。'),
        ('参考视频', '至少 1 支，标出喜欢的 3 个时间点。', '说明是喜欢“节奏、排版、人物切换”还是“音效”。'),
        ('音效与配乐', '可商用来源、文件夹结构。', '提供 WAV/MP3，并区分“章节、数字、风险、片尾”。'),
    ], [1650, 3550, 4160])

    add_heading(doc, '三、固定镜头语法：如何排版才像“在对话”', 1)
    add_image(doc, '01_hook.png', '镜头 0｜开场钩子：用问题或反差数据先抓住观众，再进入栏目片头。')
    add_table(doc, ['镜头类型', '画面规则', '推荐时长'], [
        ('钩子', '单人提问 + 1–3 个反差指标；不先解释。', '8–15 秒'),
        ('栏目片头', '钩子结束后接人物主标片头；无具体公司数据。', '2.23 秒'),
        ('提问镜头', '女主持人右侧高亮；左侧男主持人弱化；只留一行问题。', '2–6 秒'),
        ('回答镜头', '男主持人左侧高亮；右侧女主持人弱化；数据面板为视觉中心。', '6–16 秒'),
        ('章节卡', '全屏暗遮罩 + 章节编号与标题；与声音错开。', '1.2–1.6 秒'),
        ('收束', '男声说投资结论时仍是分析画面；最后女声开始时切全屏品牌片尾。', '3–4 秒（仅随末句）'),
    ], [1650, 5450, 2260])
    add_image(doc, '03_active-speaker.png', '镜头 1｜发言者高亮：当前角色放大、对方弱化；画面信息中心仍是数据面板。')
    add_callout(doc, '对话卡写法', '把整段旁白“压缩”为 8–18 字的关键句。例如：旁白讲研发费率、线控制动和验证节点；对话卡只写“Q3 主机厂定点，是兑现门槛”。', PALE_CYAN, CYAN)

    add_heading(doc, '四、章节与音画同步：这一步决定视频是否“专业”', 1)
    add_image(doc, '04_chapter-card.png', '镜头 2｜章节卡占用前后两句之间的转场空档；它退出后，人物与新面板才进入。')
    add_number(doc, '将每条 TTS 音频的真实时长写入 timeline；不要按文字字数估算。')
    add_number(doc, '给每段旁白定义“语义触发点”：讲到第一个数字再出现第一张卡；讲到第二个验证点再出现第二项。')
    add_number(doc, '章节卡放入上一段结束后的已有停顿；如果没有自然停顿，预留约 1.3 秒，并由章节扫光音效填充。')
    add_number(doc, '章节卡最后一帧必须先消失；下一条音频第一帧才出现发言者高亮卡。')
    add_number(doc, '不要对整段音频做拉伸，也不要在视频最后添加“默认收尾音”。')

    add_heading(doc, '五、音色、配乐与音效：已验证的关键决策', 1)
    add_table(doc, ['类别', '本次选择', '使用原则'], [
        ('男声 TTS', '刘飞 2.0\nzh_male_liufei_uranus_bigtts', '用于解释、判断、结论；偏沉稳。'),
        ('女声 TTS', '流畅女声 2.0\nzh_female_liuchangnv_uranus_bigtts', '用于提问、追问、收束；节奏利落。'),
        ('背景音乐', 'immersed.mp3', '约 12% 音量；开头 1.5 秒淡入；始终让对白处于最前。'),
        ('章节', 'section-tech-sweep + headline-land', '前者用于章节切换，后者用于标题落定；仅章节使用。'),
        ('数据', 'data-tick-crisp / number-counter-fast', '数字、计数、指标卡出现时轻用；不要每个元素都响。'),
        ('风险', 'risk-alert', '仅在阈值、风险提示或警示结论出现时使用一次。'),
        ('片尾', '主标动效自带扫描音效（降到约 18%）', '与最后一句并行；女声结束即结束，不用 closing-resolve。'),
    ], [1800, 3300, 4260])

    add_heading(doc, '六、片头与片尾：直接复用的资产链接', 1)
    add_asset_link(doc, '片头｜双人物主标动效', str(ROOT / 'examples/hello-world/out/financial-intro-portrait-sfx-v1.mov'), '放在开场钩子之后、正文之前；与开场数据扫描音效匹配。', '2.23 秒 / 67 帧 / 30fps')
    add_image(doc, '02_intro.png', '片头｜钩子后的栏目主标动效：无业务数据，只建立栏目识别。')
    add_asset_link(doc, '片尾｜纯主标动效', str(ROOT / 'examples/hello-world/out/financial-intro-logo-sfx-v3.mov'), '从最后女声“每期三分钟…”第一字开始播放；完成后停在最后定帧，至女声结束。', '2.23 秒 / 67 帧 / 30fps')
    add_image(doc, '05_outro.png', '片尾｜全屏无人物、无进度条、无数据面板；本片在女声结束时立即结束。')
    add_asset_link(doc, '已完成成片（参考实现）', str(ROOT / 'examples/hello-world/out/financial-report-dialogue-v3-final-cut-preview.mp4'), '本次最终无尾音预览：2分29秒，960×540。', '149.08 秒')

    add_heading(doc, '七、交付前验收清单（让 Codex 自检）', 1)
    for item in [
        '逐条核对音频时长与 Sequence 起止帧，确认没有配音/字幕/动效提前或滞后。',
        '任意时刻只突出一位发言者；对方仅保留弱化头像。',
        '对话卡只显示关键句，不显示整段旁白。',
        '右上角栏目 Logo 与人物头像、气泡卡保持安全间距。',
        '章节标题在说话前完成淡出，不能和人物卡重叠。',
        '所有数据卡、曲线和风险提示按旁白递进出现，不能一进场全亮。',
        '片头不放具体公司数据；片尾从最后一句启动，片尾结束不再加音效或空白。',
        '至少导出并人工查看：人物切换、章节卡、片尾三个关键静帧，再导出成片。',
    ]:
        add_bullet(doc, item)

    add_callout(doc, '一句话交付标准', '观众应该感觉是在看两位主持人围绕一份财报逐层追问：谁说话谁突出；说到哪里数据就出现到哪里；章节和片尾都像节目节奏的一部分，而不是后贴上的模板。', PALE_GOLD, GOLD)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build()
